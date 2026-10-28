"""Lectura de anexos oficiales y alineación de frecuencias."""
from datetime import datetime
from pathlib import Path
import json
import re
import numpy as np
import openpyxl
import pandas as pd
from .io import config, write_frame, write_json


def expand_trm(records: list[dict]) -> pd.DataFrame:
    rows = []
    for record in records:
        if record["unidad"] != "COP":
            raise ValueError("Unidad de TRM inesperada.")
        start = pd.Timestamp(record["vigenciadesde"]).normalize()
        end = pd.Timestamp(record["vigenciahasta"]).normalize()
        if end < start or (end-start).days > 7:
            raise ValueError("Intervalo de vigencia inválido en la TRM.")
        rows.extend((date, float(record["valor"])) for date in pd.date_range(start, end))
    frame = pd.DataFrame(rows, columns=["date", "trm_cop_usd"]).sort_values("date")
    if frame.date.duplicated().any():
        raise ValueError("Vigencias superpuestas en la TRM.")
    return frame.reset_index(drop=True)


def parse_ipc(path: Path) -> pd.DataFrame:
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    if sheet.cell(9,1).value != "Mes" or "2018" not in str(sheet.cell(8,25).value):
        raise ValueError("Cambió la estructura o la base del anexo IPC.")
    rows = []
    for col in range(2, sheet.max_column+1):
        year = sheet.cell(9,col).value
        if not isinstance(year,int) or not 2000 <= year <= 2100:
            continue
        for month in range(1,13):
            value = sheet.cell(9+month,col).value
            if isinstance(value,(float,int)):
                rows.append((pd.Timestamp(year,month,1),float(value)))
    return pd.DataFrame(rows,columns=["date","cpi_index"]).sort_values("date")


def parse_trade(path: Path) -> pd.DataFrame:
    sheet = openpyxl.load_workbook(path,data_only=True).worksheets[0]
    if "FOB" not in str(sheet.cell(8,2).value):
        raise ValueError("El comercio debe estar expresado en millones de USD FOB.")
    rows = [(r[0],float(r[1]),float(r[2]),float(r[3])) for r in sheet.values if isinstance(r[0],datetime)]
    frame = pd.DataFrame(rows,columns=["date","exports_usd_m","imports_usd_m","trade_balance_source_usd_m"])
    if not np.allclose(frame.exports_usd_m-frame.imports_usd_m,frame.trade_balance_source_usd_m,atol=0.005,rtol=0):
        raise ValueError("La balanza del anexo no coincide con exportaciones menos importaciones.")
    return frame.sort_values("date")


def parse_gdp(path: Path) -> pd.DataFrame:
    book = openpyxl.load_workbook(path,data_only=True)
    series = []
    for sheet_name, name in [("Cuadro 1","gdp_real_original"),("Cuadro 4","gdp_real_sa")]:
        sheet = book[sheet_name]
        candidates = [r for r in list(sheet.values)[:35] if str(r[2]).strip().lower()=="producto interno bruto"]
        if len(candidates)!=1:
            raise ValueError("No se identificó una única fila de nivel del PIB.")
        row = candidates[0]
        year, revision, data = None, None, []
        for col in range(4,sheet.max_column+1):
            header = sheet.cell(12,col).value
            if header is not None and re.match(r"^20\d{2}",str(header)):
                year = int(str(header)[:4])
                revision = "preliminar" if "pr" in str(header) else "provisional" if "p" in str(header) else "sin marca"
            quarter = sheet.cell(13,col).value
            value = row[col-1]
            if year and quarter in ["I","II","III","IV"] and isinstance(value,(int,float)):
                month = {"I":1,"II":4,"III":7,"IV":10}[quarter]
                data.append((pd.Timestamp(year,month,1),float(value),revision))
        series.append(pd.DataFrame(data,columns=["date",name,"revision"]).set_index("date"))
    frame = series[0].join(series[1].drop(columns="revision"),how="inner").reset_index()
    frame["gdp_yoy_pct"] = frame.gdp_real_original.pct_change(4,fill_method=None)*100
    frame["gdp_qoq_sa_pct"] = frame.gdp_real_sa.pct_change(fill_method=None)*100
    return frame


def parse_brent(path: Path) -> pd.DataFrame:
    frame = pd.read_excel(path,sheet_name="Data 1",header=2)
    if len(frame.columns)!=2 or "Brent" not in frame.columns[1]:
        raise ValueError("Serie Brent inesperada.")
    frame.columns = ["date","brent_usd_bbl"]
    frame["date"] = pd.to_datetime(frame.date)
    frame["brent_usd_bbl"] = pd.to_numeric(frame.brent_usd_bbl,errors="raise")
    return frame.dropna().sort_values("date")


def validate_series(frame: pd.DataFrame, frequency: str, required: list[str]) -> None:
    if frame.date.duplicated().any() or not frame.date.is_monotonic_increasing:
        raise ValueError("Fechas duplicadas o desordenadas.")
    expected = pd.date_range(frame.date.min(),frame.date.max(),freq=frequency)
    if not pd.DatetimeIndex(frame.date).equals(expected):
        raise ValueError("Hay huecos en la frecuencia de la serie.")
    if frame[required].isna().any().any():
        raise ValueError("Existen valores faltantes en la muestra de análisis.")
    if not np.isfinite(frame[required].to_numpy(dtype=float)).all():
        raise ValueError("La muestra contiene valores no finitos.")


def build(root: Path) -> None:
    settings = config(root)
    raw, out = root/"data/raw",root/"data/processed"
    start, end = pd.Timestamp(settings["monthly_start"]),pd.Timestamp(settings["monthly_end"])
    daily = expand_trm(json.loads((raw/"trm.json").read_text(encoding="utf-8")))
    daily = daily[daily.date.between(start,end+pd.offsets.MonthEnd(0))].copy()
    if len(daily)!=len(pd.date_range(start,end+pd.offsets.MonthEnd(0))):
        raise ValueError("La TRM no cubre todos los días calendario del corte.")
    fx = daily.set_index("date").resample("MS").agg(trm_cop_usd=("trm_cop_usd","mean"),fx_days=("trm_cop_usd","size"))
    brent = parse_brent(raw/"brent.xls")
    brent = brent[brent.date.between(start,end+pd.offsets.MonthEnd(0))]
    oil = brent.set_index("date").resample("MS").agg(brent_usd_bbl=("brent_usd_bbl","mean"),brent_days=("brent_usd_bbl","size"))
    ipc = parse_ipc(raw/"ipc.xlsx").set_index("date")
    trade = parse_trade(raw/"trade.xlsx").set_index("date")
    monthly = fx.join(ipc).join(trade).join(oil)
    for base, prefix in [("trm_cop_usd","trm"),("cpi_index","inflation"),("exports_usd_m","exports"),("imports_usd_m","imports"),("brent_usd_bbl","brent")]:
        monthly[prefix+"_mom_pct"] = monthly[base].pct_change(fill_method=None)*100
        monthly[prefix+"_yoy_pct"] = monthly[base].pct_change(12,fill_method=None)*100
    monthly["trade_balance_usd_m"] = monthly.exports_usd_m-monthly.imports_usd_m
    monthly["year"] = monthly.index.year
    monthly["month"] = monthly.index.month
    monthly["regime"] = np.select([monthly.trm_yoy_pct<0,monthly.trm_yoy_pct>0],["Apreciación","Depreciación"],default="Sin cambio")
    monthly["is_main_sample"] = monthly.index.to_series().between(settings["sample_start"],settings["sample_end"]).astype(int)
    monthly["is_complete"] = monthly[["trm_cop_usd","cpi_index","exports_usd_m","imports_usd_m","brent_usd_bbl"]].notna().all(axis=1).astype(int)
    monthly = monthly.reset_index()
    required = ["trm_cop_usd","inflation_yoy_pct","exports_yoy_pct","imports_yoy_pct","brent_yoy_pct"]
    main = monthly[monthly.is_main_sample==1]
    validate_series(main,"MS",required)
    if len(main)!=132:
        raise ValueError("La muestra 2015–2025 debe contener 132 meses.")
    gdp = parse_gdp(raw/"pib.xlsx")
    gdp = gdp[gdp.date>=start].copy()
    # Las medias trimestrales conservan el peso de cada día observado.
    q = monthly.set_index("date").copy()
    q["fx_weighted"] = q.trm_cop_usd*q.fx_days
    q["brent_weighted"] = q.brent_usd_bbl*q.brent_days
    agg = q.resample("QS").agg(fx_weighted=("fx_weighted","sum"),fx_days=("fx_days","sum"),brent_weighted=("brent_weighted","sum"),brent_days=("brent_days","sum"),exports_usd_m=("exports_usd_m",lambda s:s.sum(min_count=3)),imports_usd_m=("imports_usd_m",lambda s:s.sum(min_count=3)),months=("trm_cop_usd","count"))
    agg = agg[agg.months==3].copy()
    agg["trm_cop_usd"] = agg.fx_weighted/agg.fx_days
    agg["brent_usd_bbl"] = agg.brent_weighted/agg.brent_days
    quarterly = gdp.set_index("date").join(agg[["trm_cop_usd","brent_usd_bbl","exports_usd_m","imports_usd_m"]],how="inner")
    quarterly["trm_yoy_pct"] = quarterly.trm_cop_usd.pct_change(4,fill_method=None)*100
    quarterly["trm_qoq_pct"] = quarterly.trm_cop_usd.pct_change(fill_method=None)*100
    quarterly["brent_yoy_pct"] = quarterly.brent_usd_bbl.pct_change(4,fill_method=None)*100
    quarterly["year"] = quarterly.index.year
    quarterly["quarter"] = quarterly.index.quarter
    quarterly["is_main_sample"] = quarterly.index.to_series().between(settings["sample_start"],settings["sample_end"]).astype(int)
    quarterly = quarterly.reset_index()
    validate_series(quarterly[quarterly.is_main_sample==1],"QS",["gdp_yoy_pct","gdp_qoq_sa_pct","trm_yoy_pct"])
    if len(quarterly[quarterly.is_main_sample==1])!=44:
        raise ValueError("La muestra trimestral debe contener 44 observaciones.")
    calendar = pd.DataFrame({"date":pd.date_range(start,end+pd.offsets.MonthEnd(0))})
    calendar["year"] = calendar.date.dt.year
    calendar["month"] = calendar.date.dt.month
    calendar["month_label"] = calendar.date.dt.strftime("%Y-%m")
    calendar["quarter_label"] = calendar.date.dt.year.astype(str)+"-T"+calendar.date.dt.quarter.astype(str)
    for name, frame in [("trm_daily",daily),("brent_daily",brent),("monthly",monthly),("quarterly",quarterly),("calendar",calendar)]:
        write_frame(out/f"{name}.csv",frame)
    checks = {"status":"passed","monthly_observations":len(main),"quarterly_observations":44,"trm_calendar_days":len(daily),"duplicate_dates":0,"missing_main_sample":0,"trade_balance_max_error_usd_m":float((monthly.trade_balance_usd_m-monthly.trade_balance_source_usd_m).abs().max()),"monthly_coverage_end":str(monthly.date.max().date()),"common_monthly_end":str(monthly.loc[monthly.is_complete==1,"date"].max().date()),"quarterly_coverage_end":str(quarterly.date.max().date()),"imputation":"none"}
    write_json(root/"reports/data_quality.json",checks)
    print(f"Datos preparados: {len(monthly)} meses, {len(quarterly)} trimestres.")
